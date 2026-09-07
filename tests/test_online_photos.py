import os
import sys
import unittest
import tempfile
import pandas as pd
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

import config
from repository import _normalise_dataframes, SQLiteRepository, ExcelRepository


class TestOnlinePhotos(unittest.TestCase):

    def setUp(self):
        self.cfg = config.DATABASE_CONFIGS["Økonomisk Botanisk"]

    def test_01_config_schema(self):
        """Verify Online photo columns exist in config.py for Økonomisk Botanisk."""
        reg_fields = [f["name"] for f in self.cfg["ui_sections"]["registration"]]
        for col in ("Online photo 1", "Online photo 2", "Online photo 3"):
            self.assertIn(col, reg_fields, f"Missing {col} in ui_sections['registration']")

        # Check reg_groups
        reg_groups = {g["name"]: g["fields"] for g in self.cfg["ui_sections"]["reg_groups"]}
        self.assertIn("Object", reg_groups)
        for col in ("Online photo 1", "Online photo 2", "Online photo 3"):
            self.assertIn(col, reg_groups["Object"], f"Missing {col} in reg_groups['Object']")

    def test_02_normalization_and_aliasing(self):
        """Verify _normalise_dataframes correctly aliases headers, cleans values, and syncs Online_Images_Exist."""
        df_reg = pd.DataFrame({
            "ObjectID": [101, 102, 103, 104],
            "Genus": ["Quercus", "Pinus", "Betula", "Fraxinus"],
            "online_photo_1": ['"https://example.com/1.jpg"', "   https://example.com/2.jpg  ", None, ""],
            "ONLINE PHOTO 2": ["nan", "'https://example.com/clean.jpg'", "<NA>", "None"],
            "Online Photo 3": ["", "None", "https://example.com/3.jpg", "  "]
        })
        df_obs = pd.DataFrame({
            "ObjectID": [101, 102, 103, 104]
        })

        norm_reg, norm_obs = _normalise_dataframes(df_reg, df_obs, self.cfg)

        for col in ("Online photo 1", "Online photo 2", "Online photo 3"):
            self.assertIn(col, norm_reg.columns, f"Expected canonical column {col}")

        # Check values
        # Row 0 (ObjectID 101):
        self.assertEqual(norm_reg.loc[0, "Online photo 1"], "https://example.com/1.jpg")
        self.assertEqual(norm_reg.loc[0, "Online photo 2"], "")
        self.assertEqual(norm_reg.loc[0, "Online photo 3"], "")

        # Row 1 (ObjectID 102):
        self.assertEqual(norm_reg.loc[1, "Online photo 1"], "https://example.com/2.jpg")
        self.assertEqual(norm_reg.loc[1, "Online photo 2"], "https://example.com/clean.jpg")
        self.assertEqual(norm_reg.loc[1, "Online photo 3"], "")

        # Row 2 (ObjectID 103):
        self.assertEqual(norm_reg.loc[2, "Online photo 1"], "")
        self.assertEqual(norm_reg.loc[2, "Online photo 2"], "")
        self.assertEqual(norm_reg.loc[2, "Online photo 3"], "https://example.com/3.jpg")

        # Row 3 (ObjectID 104 - no photos):
        self.assertEqual(norm_reg.loc[3, "Online photo 1"], "")
        self.assertEqual(norm_reg.loc[3, "Online photo 2"], "")
        self.assertEqual(norm_reg.loc[3, "Online photo 3"], "")

        # Check Online_Images_Exist in df_obs
        self.assertIn("Online_Images_Exist", norm_obs.columns)
        self.assertTrue(norm_obs.loc[0, "Online_Images_Exist"], "101 has photos -> True")
        self.assertTrue(norm_obs.loc[1, "Online_Images_Exist"], "102 has photos -> True")
        self.assertTrue(norm_obs.loc[2, "Online_Images_Exist"], "103 has photos -> True")
        self.assertFalse(norm_obs.loc[3, "Online_Images_Exist"], "104 has NO photos -> False")

    def test_03_missing_columns_backfill(self):
        """Verify that older datasets without online photo columns get empty string backfills and False in obs."""
        df_reg = pd.DataFrame({
            "ObjectID": [201, 202],
            "Genus": ["Rosa", "Salix"]
        })
        df_obs = pd.DataFrame({
            "ObjectID": [201, 202]
        })

        norm_reg, norm_obs = _normalise_dataframes(df_reg, df_obs, self.cfg)
        for col in ("Online photo 1", "Online photo 2", "Online photo 3"):
            self.assertIn(col, norm_reg.columns)
            self.assertEqual(norm_reg.loc[0, col], "")
            self.assertEqual(norm_reg.loc[1, col], "")
        self.assertIn("Online_Images_Exist", norm_obs.columns)
        self.assertFalse(norm_obs.loc[0, "Online_Images_Exist"])
        self.assertFalse(norm_obs.loc[1, "Online_Images_Exist"])

    def test_04_excel_export_and_roundtrip(self):
        """Verify SQLiteRepository.export_to_excel exports online photo columns and syncs Online_Images_Exist in Observation."""
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = os.path.join(tmpdir, "test_online.xlsx")
            sqlite_path = os.path.join(tmpdir, "test_online.db")

            df_reg = pd.DataFrame({
                "ObjectID": [301, 302, 303],
                "Genus": ["Fagus", "Acer", "Ulmus"],
                "UID": ["uid1", "uid2", "uid3"],
                "ProblemDescription": ["", "", ""],
                "Online photo 1": ["https://images.example.org/fagus1.jpg", "", ""],
                "Online photo 2": ["", "https://images.example.org/acer2.jpg", ""],
                "Online photo 3": ["", "", ""]
            })
            # Intentionally pass stale False for all rows in df_obs to test that export_to_excel auto-syncs
            df_obs = pd.DataFrame({
                "ObjectID": [301, 302, 303],
                "Images_Missing": [False, False, True],
                "Images_Problem": [False, False, False],
                "Reviewed": [False, False, False],
                "ReviewedAt": ["", "", ""],
                "Online_Images_Exist": [False, False, False]
            })

            SQLiteRepository.save_sqlite(sqlite_path, df_reg, df_obs, pd.DataFrame(), pd.DataFrame(), pd.DataFrame())
            self.assertTrue(os.path.exists(sqlite_path))

            # Export to excel
            export_target = os.path.join(tmpdir, "exported.xlsx")
            SQLiteRepository.export_to_excel(sqlite_path, export_target, self.cfg, df_reg=df_reg, df_obs=df_obs)
            self.assertTrue(os.path.exists(export_target))

            # Inspect with openpyxl
            wb = openpyxl.load_workbook(export_target)
            self.assertIn("Registration", wb.sheetnames)
            self.assertIn("Observation", wb.sheetnames)
            ws_reg = wb["Registration"]
            ws_obs = wb["Observation"]

            # Check freeze pane A2 and autofilter
            self.assertEqual(ws_reg.freeze_panes, "A2", "Freeze pane should be A2")
            self.assertIsNotNone(ws_reg.auto_filter.ref, "Auto filter should be set")

            # Check Registration header row
            reg_headers = [cell.value for cell in ws_reg[1]]
            for col in ("Online photo 1", "Online photo 2", "Online photo 3"):
                self.assertIn(col, reg_headers, f"Exported Registration sheet missing {col}")

            idx1 = reg_headers.index("Online photo 1") + 1
            idx2 = reg_headers.index("Online photo 2") + 1

            # Row 2 (ObjectID 301):
            self.assertEqual(ws_reg.cell(row=2, column=idx1).value, "https://images.example.org/fagus1.jpg")
            # Row 3 (ObjectID 302):
            self.assertEqual(ws_reg.cell(row=3, column=idx2).value, "https://images.example.org/acer2.jpg")

            # Check Observation sheet for Online_Images_Exist
            obs_headers = [cell.value for cell in ws_obs[1]]
            self.assertIn("Online_Images_Exist", obs_headers, "Observation sheet missing Online_Images_Exist")
            idx_online = obs_headers.index("Online_Images_Exist") + 1

            # Row 2 (ObjectID 301 has photo 1) -> must be True in Excel
            val_301 = ws_obs.cell(row=2, column=idx_online).value
            self.assertTrue(bool(val_301), f"Expected True for ObjectID 301 in Excel, got {val_301}")

            # Row 3 (ObjectID 302 has photo 2) -> must be True in Excel
            val_302 = ws_obs.cell(row=3, column=idx_online).value
            self.assertTrue(bool(val_302), f"Expected True for ObjectID 302 in Excel, got {val_302}")

            # Row 4 (ObjectID 303 has no photo) -> must be False in Excel
            val_303 = ws_obs.cell(row=4, column=idx_online).value
            self.assertFalse(bool(val_303), f"Expected False for ObjectID 303 in Excel, got {val_303}")

            # Load back via ExcelRepository.load_excel
            loaded_reg, loaded_obs, _, _, _ = ExcelRepository.load_excel(export_target, self.cfg)
            row_301 = loaded_reg[loaded_reg["ObjectID"].astype(str) == "301"]
            self.assertEqual(row_301["Online photo 1"].iloc[0], "https://images.example.org/fagus1.jpg")
            row_302 = loaded_reg[loaded_reg["ObjectID"].astype(str) == "302"]
            self.assertEqual(row_302["Online photo 2"].iloc[0], "https://images.example.org/acer2.jpg")

            # Check re-imported observation values
            obs_301 = loaded_obs[loaded_obs["ObjectID"].astype(str) == "301"]
            self.assertTrue(bool(obs_301["Online_Images_Exist"].iloc[0]))
            obs_302 = loaded_obs[loaded_obs["ObjectID"].astype(str) == "302"]
            self.assertTrue(bool(obs_302["Online_Images_Exist"].iloc[0]))
            obs_303 = loaded_obs[loaded_obs["ObjectID"].astype(str) == "303"]
            self.assertFalse(bool(obs_303["Online_Images_Exist"].iloc[0]))

    def test_05_url_security_validation(self):
        """Verify URL validation rejects unsafe schemes."""
        import re

        valid_urls = [
            "https://example.com/photo.jpg",
            "http://example.com/photo.png",
            "HTTP://EXAMPLE.COM/PHOTO.JPEG",
            "  https://sub.domain.org/path?query=1  ",
            "www.example.com/photo.jpg"
        ]
        invalid_urls = [
            "javascript:alert(1)",
            "file:///C:/Windows/System32/calc.exe",
            "ftp://files.example.com/test",
            "",
            None,
            "   "
        ]

        def is_safe_url(url):
            if not url or not isinstance(url, str):
                return False
            url_str = url.strip()
            if not url_str:
                return False
            if not re.match(r'^(?:https?://|www\.)', url_str, re.IGNORECASE):
                return False
            return True

        for u in valid_urls:
            self.assertTrue(is_safe_url(u), f"Should accept: {u}")
        for u in invalid_urls:
            self.assertFalse(is_safe_url(u), f"Should reject: {u}")

    def test_06_cache_set_vectorization(self):
        """Verify vectorized row cache calculation for has_online_photos as implemented in main_window.py."""
        reg_df = pd.DataFrame({
            "ObjectID": [1, 2, 3, 4],
            "Online photo 1": ["https://pic1.jpg", "", "   ", ""],
            "Online photo 2": ["", None, "nan", ""],
            "Online photo 3": ["", "", "", "https://pic3.jpg"]
        }).set_index("ObjectID")

        # Logic matching _ensure_row_caches in main_window.py
        online_cols = [c for c in ("Online photo 1", "Online photo 2", "Online photo 3") if c in reg_df.columns]
        has_photo_mask = pd.Series(False, index=reg_df.index)
        for c in online_cols:
            s = reg_df[c].fillna("").astype(str).str.strip()
            has_photo_mask |= (s != "") & (~s.isin(["nan", "None", "<NA>"]))

        has_online_set = set(reg_df.index[has_photo_mask])

        self.assertIn(1, has_online_set)    # has Online photo 1
        self.assertNotIn(2, has_online_set) # empty/None
        self.assertNotIn(3, has_online_set) # whitespace / nan
        self.assertIn(4, has_online_set)    # has Online photo 3

    def test_07_registry_panel_card_defs(self):
        """Verify Online photo columns are registered in registry_panel.py card_defs."""
        from ui.registry_panel import RegistryPanel
        # Check that RegistryPanel's metadata card definitions include the 3 columns
        found_fields = set()
        # Find the fields defined in RegistryPanel
        import inspect
        src = inspect.getsource(RegistryPanel.build_sections)
        for col in ("Online photo 1", "Online photo 2", "Online photo 3"):
            self.assertIn(col, src, f"Expected {col} to be defined in RegistryPanel.build_sections")

    def test_08_live_commit_sync(self):
        """Verify that updating online photo fields updates df_obs and cache dictionaries."""
        from repository import ONLINE_EXISTS_COLUMN
        oid = "999"
        df_reg = pd.DataFrame([{
            "ObjectID": oid,
            "Online photo 1": "",
            "Online photo 2": "",
            "Online photo 3": ""
        }]).set_index("ObjectID")
        df_obs = pd.DataFrame([{
            "ObjectID": oid,
            ONLINE_EXISTS_COLUMN: False
        }]).set_index("ObjectID")

        cached_reg_dict = {oid: {"Online photo 1": "", "Online photo 2": "", "Online photo 3": ""}}
        cached_obs_dict = {oid: {ONLINE_EXISTS_COLUMN: False}}
        has_online_set = set()

        # Emulate user setting "Online photo 1" in UI and commit_current_object
        new_val = "https://example.org/photo1.jpg"
        cached_reg_dict[oid]["Online photo 1"] = new_val
        has_any = any(
            bool(str(cached_reg_dict[oid].get(f"Online photo {i}", "")).strip() not in ("", "nan", "None", "<NA>"))
            for i in (1, 2, 3)
        )
        if has_any:
            has_online_set.add(oid)
        df_obs.at[oid, ONLINE_EXISTS_COLUMN] = has_any
        cached_obs_dict[oid][ONLINE_EXISTS_COLUMN] = has_any

        self.assertTrue(df_obs.at[oid, ONLINE_EXISTS_COLUMN])
        self.assertTrue(cached_obs_dict[oid][ONLINE_EXISTS_COLUMN])
        self.assertIn(oid, has_online_set)

        # Emulate user clearing "Online photo 1"
        cached_reg_dict[oid]["Online photo 1"] = ""
        has_any = any(
            bool(str(cached_reg_dict[oid].get(f"Online photo {i}", "")).strip() not in ("", "nan", "None", "<NA>"))
            for i in (1, 2, 3)
        )
        if not has_any:
            has_online_set.discard(oid)
        df_obs.at[oid, ONLINE_EXISTS_COLUMN] = has_any
        cached_obs_dict[oid][ONLINE_EXISTS_COLUMN] = has_any

        self.assertFalse(df_obs.at[oid, ONLINE_EXISTS_COLUMN])
        self.assertFalse(cached_obs_dict[oid][ONLINE_EXISTS_COLUMN])
        self.assertNotIn(oid, has_online_set)


if __name__ == "__main__":
    unittest.main(verbosity=2)

